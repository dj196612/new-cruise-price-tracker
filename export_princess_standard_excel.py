from __future__ import annotations

import asyncio
import json
import os
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from playwright.async_api import async_playwright

import collectors.princess as base
import collectors.princess_runtime as runtime


STANDARD_FARE_TYPE = "BESTFARE"
EXPECTED_CRUISES = 93
ALLOWED_PORTS = {"sydney", "brisbane", "melbourne"}
CABINS = ["interior", "oceanview", "balcony", "mini_suite", "suite"]
CABIN_LABELS = {
    "interior": "Interior",
    "oceanview": "Oceanview",
    "balcony": "Balcony",
    "mini_suite": "Mini-Suite",
    "suite": "Suite",
}

# Force Princess's pricing API parser to use BESTFARE (Standard) rather than
# BESTVALUE (Princess Plus).  The existing tracker intentionally uses Plus for
# its live dashboard, so this export remains isolated from that behaviour.
runtime.SOLO_FARE_TYPE = STANDARD_FARE_TYPE
_original_cabin_prices = runtime._solo_cabin_prices


def _standard_cabin_prices(cruises, fare_type=None, mode=None):
    return _original_cabin_prices(
        cruises,
        fare_type=fare_type or STANDARD_FARE_TYPE,
        mode=mode,
    )


runtime._solo_cabin_prices = _standard_cabin_prices
runtime._price_mode["value"] = "fare"
runtime._price_mode["scale"] = 1.0


def _normalise_port(value: str | None) -> str:
    raw = (value or "").strip().lower()
    if raw.startswith("sydney") or raw == "syd":
        return "Sydney"
    if raw.startswith("brisbane") or raw == "bne":
        return "Brisbane"
    if raw.startswith("melbourne") or raw == "mel":
        return "Melbourne"
    return (value or "").strip()


def _in_scope(record: dict) -> bool:
    return _normalise_port(record.get("departure_port")).lower() in ALLOWED_PORTS


async def _collect_catalogue(headless: bool = True) -> list[dict]:
    """Collect Princess's 93 SYD/BNE/MEL sailings and metadata.

    The rendered Princess page currently shows Plus prices. We use it only for
    discovery/metadata, then replace every price from the BESTFARE API.
    """
    original_set_guest = base.set_guest_count_on_detail
    original_scrape_detail = base.scrape_detail
    base.set_guest_count_on_detail = runtime._disable_page_solo
    base.scrape_detail = runtime._scrape_detail_fixed
    try:
        records = await base.collect(headless=headless)
    finally:
        base.set_guest_count_on_detail = original_set_guest
        base.scrape_detail = original_scrape_detail

    runtime._validate_metadata(records)

    scoped = [r for r in records if _in_scope(r)]
    for record in scoped:
        record["departure_port"] = _normalise_port(record.get("departure_port"))

    print(f"Princess catalogue: {len(records)} total records, {len(scoped)} from Sydney/Brisbane/Melbourne")
    if len(scoped) != EXPECTED_CRUISES:
        raise RuntimeError(
            f"Expected {EXPECTED_CRUISES} Princess cruises from Sydney/Brisbane/Melbourne, "
            f"but collected {len(scoped)}. Refusing to export an incomplete workbook."
        )
    return scoped


async def _price_all_standard(records: list[dict], headless: bool = True) -> None:
    seed_url = next((r.get("url") for r in records if r.get("url")), None)
    if not seed_url:
        raise RuntimeError("No Princess itinerary URL available to seed pricing session.")

    async with async_playwright() as p:
        browser = await p.chromium.launch(headless=headless)
        context = await browser.new_context(
            locale="en-AU",
            timezone_id="Australia/Sydney",
            viewport={"width": 1440, "height": 1000},
            extra_http_headers={"Accept-Language": "en-AU,en;q=0.9"},
        )
        await context.add_cookies([
            {"name": "cookie_geo_country", "value": "AU", "domain": ".princess.com", "path": "/"},
            {"name": "cookie_url_country", "value": "AU", "domain": ".princess.com", "path": "/"},
        ])

        async def force_locale(route):
            await route.fulfill(status=200, content_type="application/json", body=json.dumps(base.AU_LOCALE))

        await context.route(base.LOCALE_URL, force_locale)
        headers, template_payload = await runtime._capture_pricing_session(context, seed_url)
        session = runtime._PricingSession(context, seed_url, headers, template_payload)
        session.two_guest_ok = True

        semaphore = asyncio.Semaphore(3)

        async def price_record(index: int, record: dict):
            code = runtime._voyage_code(record)
            if not code:
                record["pricing_note"] = "Missing voyage code"
                return

            async with semaphore:
                notes = []
                for guests in (1, 2):
                    prices, statuses, cruises, data, fare_notes = await runtime._fetch_fares(session, code, guests)
                    notes.extend(f"{guests}g: {n}" for n in fare_notes)

                    if guests == 1:
                        for cabin in CABINS:
                            record[f"one_guest_{cabin}"] = prices.get(cabin) if prices else None
                        record["one_guest_price"] = record.get("one_guest_interior")
                    else:
                        for cabin in CABINS:
                            record[f"two_guest_{cabin}_pp"] = prices.get(cabin) if prices else None
                        record["two_guest_pp_price"] = record.get("two_guest_interior_pp")
                        record["two_guest_total"] = (
                            round(record["two_guest_pp_price"] * 2, 2)
                            if record.get("two_guest_pp_price") is not None else None
                        )

                    if not prices or not any(v is not None for v in prices.values()):
                        notes.append(
                            f"{guests}g: no bookable Standard fare; statuses={dict(statuses)}"
                        )

                    if not record.get("departure_date") and cruises:
                        record["departure_date"] = runtime._sail_date(cruises) or record.get("departure_date")

                record["currency"] = "AUD"
                record["fare_basis"] = "Princess Standard (BESTFARE)"
                record["pricing_note"] = "; ".join(notes)
                print(
                    f"Standard {index}/{len(records)} {code}: "
                    f"1g Interior={record.get('one_guest_interior')} | "
                    f"2g Interior pp={record.get('two_guest_interior_pp')}"
                )

        await asyncio.gather(*(price_record(i, record) for i, record in enumerate(records, 1)))
        await browser.close()

    if len(records) != EXPECTED_CRUISES:
        raise RuntimeError(f"Row count changed unexpectedly: {len(records)}")

    one_rows = sum(1 for r in records if any(r.get(f"one_guest_{c}") is not None for c in CABINS))
    two_rows = sum(1 for r in records if any(r.get(f"two_guest_{c}_pp") is not None for c in CABINS))
    print(f"Standard pricing captured: 1 guest {one_rows}/{len(records)}, 2 guests {two_rows}/{len(records)}")


def _style_sheet(ws, freeze="A2"):
    ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top")

    for col_idx, column_cells in enumerate(ws.columns, 1):
        max_len = 0
        for cell in column_cells[:100]:
            value = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, len(value))
        width = min(max(max_len + 2, 11), 42)
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def _add_link(cell, url: str | None):
    if url:
        cell.value = "Open Princess"
        cell.hyperlink = url
        cell.style = "Hyperlink"


def _write_guest_sheet(wb: Workbook, records: list[dict], guests: int):
    title = "1 Guest" if guests == 1 else "2 Guests"
    ws = wb.create_sheet(title)

    if guests == 1:
        headers = [
            "Departure Date", "Nights", "Ship", "Cruise", "Departure Port", "Voyage Code",
            "Interior Total", "Oceanview Total", "Balcony Total", "Mini-Suite Total", "Suite Total",
            "Currency", "Fare Basis", "Direct Princess Link", "Pricing Note",
        ]
    else:
        headers = [
            "Departure Date", "Nights", "Ship", "Cruise", "Departure Port", "Voyage Code",
            "Interior pp", "Interior 2-Guest Total", "Oceanview pp", "Oceanview 2-Guest Total",
            "Balcony pp", "Balcony 2-Guest Total", "Mini-Suite pp", "Mini-Suite 2-Guest Total",
            "Suite pp", "Suite 2-Guest Total", "Currency", "Fare Basis", "Direct Princess Link", "Pricing Note",
        ]
    ws.append(headers)

    for r in records:
        if guests == 1:
            row = [
                r.get("departure_date"), r.get("nights"), r.get("ship"), r.get("name"),
                r.get("departure_port"), runtime._voyage_code(r),
                r.get("one_guest_interior"), r.get("one_guest_oceanview"), r.get("one_guest_balcony"),
                r.get("one_guest_mini_suite"), r.get("one_guest_suite"),
                "AUD", "Princess Standard", None, r.get("pricing_note") or "",
            ]
        else:
            def pp(c):
                return r.get(f"two_guest_{c}_pp")

            def total(c):
                value = pp(c)
                return round(value * 2, 2) if value is not None else None

            row = [
                r.get("departure_date"), r.get("nights"), r.get("ship"), r.get("name"),
                r.get("departure_port"), runtime._voyage_code(r),
                pp("interior"), total("interior"), pp("oceanview"), total("oceanview"),
                pp("balcony"), total("balcony"), pp("mini_suite"), total("mini_suite"),
                pp("suite"), total("suite"), "AUD", "Princess Standard", None,
                r.get("pricing_note") or "",
            ]
        ws.append(row)
        _add_link(ws.cell(ws.max_row, headers.index("Direct Princess Link") + 1), r.get("url"))

    money_columns = []
    for idx, h in enumerate(headers, 1):
        if "Interior" in h or "Oceanview" in h or "Balcony" in h or "Mini-Suite" in h or "Suite" in h:
            money_columns.append(idx)
    for col in money_columns:
        for cell in ws.iter_cols(min_col=col, max_col=col, min_row=2, max_row=ws.max_row):
            for c in cell:
                c.number_format = '$#,##0.00'

    _style_sheet(ws)


def _write_summary(wb: Workbook, records: list[dict]):
    ws = wb.active
    ws.title = "Summary"
    priced_one = sum(1 for r in records if any(r.get(f"one_guest_{c}") is not None for c in CABINS))
    priced_two = sum(1 for r in records if any(r.get(f"two_guest_{c}_pp") is not None for c in CABINS))

    rows = [
        ["Princess Standard Fare Export", ""],
        ["Generated", datetime.now().astimezone().strftime("%Y-%m-%d %H:%M %Z")],
        ["Departure ports", "Sydney, Brisbane, Melbourne"],
        ["Fare basis", "Princess Standard (BESTFARE)"],
        ["Currency", "AUD"],
        ["Cruise rows - 1 Guest", len(records)],
        ["Cruise rows - 2 Guests", len(records)],
        ["Expected cruise rows", EXPECTED_CRUISES],
        ["1 Guest rows with at least one bookable cabin fare", priced_one],
        ["2 Guest rows with at least one bookable cabin fare", priced_two],
        ["Price convention", "1 Guest = total fare for one guest; 2 Guests = per-person fare plus calculated two-person total"],
        ["Verification", "Use the direct Princess link on each row. Unavailable cabin categories are left blank."],
    ]
    for row in rows:
        ws.append(row)
    ws["A1"].font = Font(bold=True, size=16)
    ws.column_dimensions["A"].width = 48
    ws.column_dimensions["B"].width = 90
    for cell in ws["A"]:
        cell.font = Font(bold=True)


def _build_workbook(records: list[dict], output: Path):
    records = sorted(
        records,
        key=lambda r: (r.get("departure_date") or "9999-99-99", r.get("departure_port") or "", runtime._voyage_code(r) or ""),
    )
    wb = Workbook()
    _write_summary(wb, records)
    _write_guest_sheet(wb, records, 1)
    _write_guest_sheet(wb, records, 2)
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    print(f"Wrote {output} with {len(records)} rows on each guest sheet")


async def main():
    records = await _collect_catalogue(headless=True)
    await _price_all_standard(records, headless=True)
    output = Path("reports/Princess_Standard_Fares.xlsx")
    _build_workbook(records, output)


if __name__ == "__main__":
    asyncio.run(main())
